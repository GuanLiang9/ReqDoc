"""Register -> Markdown and Word. Both are built from the same section data so they never differ."""
import io
import re

from . import config
from .models import Register

EMPTY = "Not discussed in this meeting"
WARN = "⚠ to confirm"
TEMPLATE = config.ROOT / "templates" / "requirement_template.docx"

# A block is one of:
#   ("empty",) ("text", str) ("sub", str) ("bullets", [str]) ("numbered", [str])
#   ("table", [headers], [[cells]]) ("mermaid", str)
Block = tuple


def _flag(item) -> str:
    return f" {WARN}" if item.needs_confirmation else ""


def _ts(item) -> str:
    return f"[{item.timestamp}]" if item.timestamp else ""


def _bullets(items: list[str]) -> Block:
    return ("bullets", items) if items else ("empty",)


def _numbered(items: list[str]) -> Block:
    return ("numbered", items) if items else ("empty",)


def _table(headers: list[str], rows: list[list[str]]) -> Block:
    return ("table", headers, rows) if rows else ("empty",)


def sections(reg: Register) -> list[tuple[str, list[Block]]]:
    """The 18 fixed sections of the requirement document."""
    meeting = reg.meetings[-1] if reg.meetings else None
    fr = [r for r in reg.requirements if r.kind == "FR"]
    nfr = [r for r in reg.requirements if r.kind == "NFR"]
    req_headers = ["ID", "Requirement", "Priority", "Requested by", "Time", "Status"]

    def req_rows(items):
        return [[r.id, r.text + _flag(r), r.priority, r.requested_by, _ts(r), r.status] for r in items]

    if reg.diagram_mermaid:
        diagram: list[Block] = [("mermaid", reg.diagram_mermaid)]
    else:
        diagram = [("text", "Diagram not generated yet")] if reg.tobe_steps else [("empty",)]

    return [
        ("Document control", [_table(["Field", "Value"], [
            ["Project", f"{reg.project_code} {reg.project_name}".strip()],
            ["Meeting", meeting.title if meeting else ""],
            ["Date", meeting.date if meeting else ""],
            ["Attendees", ", ".join(meeting.attendees) if meeting else ""],
            ["Version", reg.version],
            ["Status", reg.doc_status],
            ["Reviewer", reg.reviewer or "TBC"],
        ])]),
        ("Meeting summary", [("text", reg.summary)] if reg.summary else [("empty",)]),
        ("Business objectives", [_bullets(reg.objectives)]),
        ("Scope", [("sub", "In scope"), _bullets(reg.scope_in), ("sub", "Out of scope"), _bullets(reg.scope_out),
                   ("sub", "Deferred"), _bullets(reg.scope_deferred)]),
        ("Stakeholders and users", [_bullets(reg.stakeholders)]),
        ("Functional requirements", [_table(req_headers, req_rows(fr))]),
        ("Non-functional requirements", [_table(req_headers, req_rows(nfr))]),
        ("Current process (as-is)", [_numbered(reg.asis_steps)]),
        ("Proposed process (to-be)", [_numbered(reg.tobe_steps)]),
        ("Workflow diagram", diagram),
        ("Business rules", [_table(["ID", "Rule", "Time"], [[r.id, r.text + _flag(r), _ts(r)] for r in reg.business_rules])]),
        ("Data and fields", [_bullets(reg.data_fields)]),
        ("Assumptions and constraints", [_bullets(reg.assumptions)]),
        ("Decisions made", [_table(["ID", "Decision", "Decided by", "Time"],
                                   [[d.id, d.text + _flag(d), d.decided_by, _ts(d)] for d in reg.decisions])]),
        ("Open questions", [_table(["ID", "Question", "Owner", "Time", "Status"],
                                   [[q.id, q.text + _flag(q), q.owner, _ts(q), q.status] for q in reg.questions])]),
        ("Action items", [_table(["ID", "Task", "Owner", "Due", "Time", "Status"],
                                 [[a.id, a.task + _flag(a), a.owner, a.due, _ts(a), a.status] for a in reg.actions])]),
        ("Risks", [_table(["ID", "Risk", "Time"], [[r.id, r.text + _flag(r), _ts(r)] for r in reg.risks])]),
        ("Glossary", [_table(["Term", "Meaning"], [[g.term, g.meaning] for g in reg.glossary])]),
    ]


def status_line(reg: Register) -> str:
    return "DRAFT — AI-generated, pending human review" if reg.doc_status == "DRAFT" else f"APPROVED by {reg.reviewer or 'reviewer'}"


def doc_title(reg: Register) -> str:
    return f"{reg.project_code} {reg.project_name} — Requirements".replace("  ", " ")


def appendix_lines(reg: Register) -> tuple[list[str], list[str]]:
    meetings = [f"{m.date} — {m.title}" + (f" ({m.source_file})" if m.source_file else "") for m in reg.meetings]
    return meetings, reg.change_log


# --- Markdown ------------------------------------------------------------------

def _md_cell(text: str) -> str:
    return (text or "").replace("|", "\\|").replace("\n", " ").strip() or "—"


def _md_block(block: Block) -> str:
    kind = block[0]
    if kind == "empty":
        return f"_{EMPTY}_"
    if kind == "text":
        return block[1]
    if kind == "sub":
        return f"**{block[1]}**"
    if kind == "bullets":
        return "\n".join(f"- {i}" for i in block[1])
    if kind == "numbered":
        return "\n".join(f"{n}. {i}" for n, i in enumerate(block[1], 1))
    if kind == "mermaid":
        return f"```mermaid\n{block[1]}\n```"
    headers, rows = block[1], block[2]
    lines = ["| " + " | ".join(headers) + " |", "|" + "---|" * len(headers)]
    lines += ["| " + " | ".join(_md_cell(c) for c in row) + " |" for row in rows]
    return "\n".join(lines)


def render_markdown(reg: Register) -> str:
    out = [f"# {doc_title(reg)}", "", f"**Status: {status_line(reg)}**", ""]
    for n, (heading, blocks) in enumerate(sections(reg), 1):
        out += [f"## {n}. {heading}", ""]
        for b in blocks:
            out += [_md_block(b), ""]
    meetings, log = appendix_lines(reg)
    out += ["## Appendix", "", "**Source meetings**", ""]
    out += [f"- {m}" for m in meetings] or ["—"]
    if log:
        out += ["", "**Change log**", ""] + [f"- {line}" for line in log]
    return "\n".join(out) + "\n"


# --- Word ----------------------------------------------------------------------

_BAD_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _clean(text: str) -> str:
    return _BAD_XML.sub("", text or "")


def _add_page_number(paragraph) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def render_docx(reg: Register) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    doc = Document(str(TEMPLATE)) if TEMPLATE.exists() else Document()
    style_names = {s.name for s in doc.styles}

    def style(name: str, fallback: str = "Normal") -> str:
        return name if name in style_names else fallback

    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    # header / footer
    sec = doc.sections[0]
    sec.header.paragraphs[0].text = _clean(f"{reg.project_code} {reg.project_name} — Requirements")
    foot = sec.footer.paragraphs[0]
    foot.text = _clean(f"v{reg.version} · {reg.doc_status} · Page ")
    _add_page_number(foot)
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(_clean(doc_title(reg)), style=style("Title"))
    banner = doc.add_paragraph()
    run = banner.add_run(_clean(status_line(reg)))
    run.bold = True
    run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B) if reg.doc_status == "DRAFT" else RGBColor(0x1E, 0x84, 0x49)

    def add_empty():
        p = doc.add_paragraph()
        p.add_run(EMPTY).italic = True

    def add_table(headers, rows):
        table = doc.add_table(rows=1, cols=len(headers))
        try:
            table.style = style("Light Grid Accent 1", "Table Grid")
        except KeyError:
            pass
        for cell, h in zip(table.rows[0].cells, headers):
            cell.text = ""
            r = cell.paragraphs[0].add_run(_clean(h))
            r.bold = True
        for row in rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = _clean(value) or "—"
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9)
        doc.add_paragraph()

    for n, (heading, blocks) in enumerate(sections(reg), 1):
        doc.add_heading(_clean(f"{n}. {heading}"), level=1)
        for b in blocks:
            kind = b[0]
            if kind == "empty":
                add_empty()
            elif kind == "text":
                doc.add_paragraph(_clean(b[1]))
            elif kind == "sub":
                doc.add_paragraph().add_run(_clean(b[1])).bold = True
            elif kind == "bullets":
                for i in b[1]:
                    doc.add_paragraph(_clean(i), style=style("List Bullet"))
            elif kind == "numbered":
                for i in b[1]:
                    doc.add_paragraph(_clean(i), style=style("List Number"))
            elif kind == "mermaid":
                doc.add_paragraph().add_run("Diagram source (Mermaid). Paste it into mermaid.live to see the picture.").italic = True
                for line in b[1].splitlines():
                    p = doc.add_paragraph()
                    r = p.add_run(_clean(line))
                    r.font.name = "Consolas"
                    r.font.size = Pt(9)
                    p.paragraph_format.space_after = Pt(0)
            else:
                add_table(b[1], b[2])

    meetings, log = appendix_lines(reg)
    doc.add_heading("Appendix", level=1)
    doc.add_paragraph().add_run("Source meetings").bold = True
    for m in meetings or ["—"]:
        doc.add_paragraph(_clean(m), style=style("List Bullet"))
    if log:
        doc.add_paragraph().add_run("Change log").bold = True
        for line in log:
            doc.add_paragraph(_clean(line), style=style("List Bullet"))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
